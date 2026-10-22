"""xlsx 读写（openpyxl）。格式判断在 rules/plan_import.py，这里只管字节 ↔ 单元格。"""
import io

from openpyxl import Workbook, load_workbook

from rules.plan_import import COLUMNS, T_DEMAND

MAX_BYTES = 5 * 1024 * 1024  # 设计取值 · 未实测（spec §2）；调用方在读之前比较


class NotXlsx(Exception):
    pass


def read_first_sheet(data: bytes) -> tuple[list[object], list[list[object]]]:
    # 任何解析失败（坏 zip、XML 损坏、包缺件…）都算「不是可读的 xlsx」；原因类型留在消息里供日志
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            it = wb.worksheets[0].iter_rows(values_only=True)
            header = list(next(it, ()))
            width = len(header)
            rows = [(list(r) + [None] * width)[:max(width, len(r))] for r in it]
        finally:
            wb.close()
    except Exception as e:
        raise NotXlsx(f"不是可读的 xlsx：{type(e).__name__}: {e}") from e
    return header, rows


def build_template(periods: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "销售计划"
    ws.append(list(COLUMNS) + list(periods))
    # 类型写成非法值：不删示例行就导入会被 bad_type 拒，示例不会当真数入库
    ws.append([f"{T_DEMAND}（示例，删掉这行）", "SKU-001", "1234", "示例店", "US", "MSKU-001"]
              + [100] * len(periods))
    ws.append(["采购（示例，删掉这行）", "SKU-001", None, None, None, None] + [80] * len(periods))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
